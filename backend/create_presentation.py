from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = Path(__file__).parent.parent / "CyberTrace_AI_Architecture_and_Demo.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(19, 38, 43)
GREEN = RGBColor(22, 132, 92)
MINT = RGBColor(216, 240, 223)
PAPER = RGBColor(244, 247, 243)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(104, 124, 126)
ORANGE = RGBColor(216, 137, 87)
RED = RGBColor(191, 98, 79)


def box(slide, x, y, w, h, fill=WHITE, line=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def text(slide, value, x, y, w, h, size=18, color=NAVY, bold=False, font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def bullets(slide, items, x, y, w, h, size=17, color=NAVY):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = "- " + item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(10)
    return shape


def base(title, eyebrow, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY if dark else PAPER
    text(slide, eyebrow.upper(), 0.65, 0.38, 4, 0.25, 9, MINT if dark else GREEN, True)
    text(slide, title, 0.65, 0.72, 12, 0.55, 28, WHITE if dark else NAVY, True, "Aptos Display")
    box(slide, 0.65, 1.42, 12.05, 0.015, MINT if dark else RGBColor(220, 228, 224))
    text(slide, f"CYBERTRACE AI  /  {len(prs.slides):02d}", 11.1, 7.1, 1.6, 0.2, 8, RGBColor(140, 170, 162) if dark else MUTED, True, align=PP_ALIGN.RIGHT)
    return slide


def code_block(slide, code, x, y, w, h):
    box(slide, x, y, w, h, NAVY, NAVY, True)
    text(slide, code, x + 0.18, y + 0.15, w - 0.36, h - 0.3, 12, MINT, False, "Consolas")

# 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
text(slide, "CYBERTRACE", 0.75, 0.8, 6, 0.45, 18, MINT, True)
text(slide, "AI", 6.35, 0.8, 1, 0.45, 18, RGBColor(139, 214, 165), True)
text(slide, "Intelligent data recovery\nand digital evidence reconstruction", 0.75, 1.85, 10.8, 1.45, 38, WHITE, True, "Aptos Display")
text(slide, "Working web application for identifying, reconstructing, classifying and prioritizing recoverable information from damaged storage data.", 0.8, 3.65, 8.8, 0.7, 18, RGBColor(190, 210, 204))
box(slide, 0.8, 5.55, 2.35, 0.7, GREEN, GREEN, True)
text(slide, "HACKATHON BUILD", 0.8, 5.55, 2.35, 0.7, 12, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
text(slide, "Cybersecurity / AI / Digital Forensics", 0.8, 6.55, 5, 0.3, 11, RGBColor(140, 170, 162))

# 2
slide = base("The challenge and the product response", "01 / Problem")
box(slide, 0.7, 1.8, 5.7, 4.65, WHITE, RGBColor(220,228,224), True)
text(slide, "THE PROBLEM", 1.0, 2.1, 2, 0.3, 10, RED, True)
bullets(slide, ["Deleted and corrupted storage may contain only partial file evidence.", "Traditional recovery returns files but not confidence, relationships or investigative context.", "Investigators need to know what is realistically restorable."], 1.0, 2.6, 4.9, 2.4, 16)
box(slide, 6.9, 1.8, 5.7, 4.65, MINT, MINT, True)
text(slide, "CYBERTRACE RESPONSE", 7.2, 2.1, 3, 0.3, 10, GREEN, True)
bullets(slide, ["Finds known artifacts and damaged clusters.", "Builds fragment relationships and best-effort reconstructions.", "Scores integrity, applies ML ranking and explains recovery decisions.", "Stores provenance and exports raw plus reconstructed evidence."], 7.2, 2.6, 4.9, 2.8, 16)

# 3
slide = base("System architecture", "02 / Architecture", True)
labels = [("WEB APPLICATION", 5.1, 1.75, 3.1, MINT), ("ANALYSIS ENGINE", 5.1, 2.65, 3.1, RGBColor(167,224,179)), ("FRAGMENT ML", 1.0, 4.1, 2.2, WHITE), ("RECONSTRUCTION", 3.55, 4.1, 2.3, WHITE), ("INTEGRITY", 6.15, 4.1, 1.9, WHITE), ("ARTIFACT ANALYSIS", 8.35, 4.1, 2.5, WHITE), ("EVIDENCE INTELLIGENCE", 4.25, 5.35, 3.3, MINT), ("REPORT / EXPORT", 4.25, 6.2, 3.3, RGBColor(167,224,179))]
for label, x, y, w, fill in labels:
    box(slide, x, y, w, 0.55, fill, fill, True)
    text(slide, label, x, y, w, 0.55, 11, NAVY, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
for x1,y1,x2,y2 in [(6.65,2.3,6.65,2.65),(6.65,3.2,2.1,4.1),(6.65,3.2,4.7,4.1),(6.65,3.2,7.1,4.1),(6.65,3.2,9.55,4.1),(6.65,4.65,5.9,5.35),(6.65,5.9,5.9,6.2)]:
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); line.line.color.rgb = RGBColor(126, 190, 157); line.line.width = Pt(1.5)
text(slide, "Read-only evidence input -> analysis -> explainable recovery output", 1.0, 6.85, 6, 0.25, 11, RGBColor(160,190,181))

# 4
slide = base("Frontend: a forensic workspace, not a static demo", "03 / Frontend")
box(slide, 0.75, 1.75, 4.1, 4.9, NAVY, NAVY, True)
text(slide, "CyberTrace dashboard", 1.05, 2.05, 3.5, 0.35, 18, WHITE, True, "Aptos Display")
for i, label in enumerate(["Dashboard", "New analysis", "Evidence", "Recovered", "Relationships", "Reports"]):
    box(slide, 1.05, 2.65 + i * 0.47, 2.9, 0.32, RGBColor(39,72,76) if i == 0 else NAVY, RGBColor(39,72,76) if i == 0 else NAVY, True)
    text(slide, label, 1.25, 2.65 + i * 0.47, 2.5, 0.32, 11, MINT if i == 0 else RGBColor(181,201,195), False, valign=MSO_ANCHOR.MIDDLE)
text(slide, "Actual frontend responsibilities", 5.35, 1.85, 4.5, 0.35, 21, GREEN, True, "Aptos Display")
bullets(slide, ["index.html: intake, progress, evidence queue, inspector, graph and report controls.", "app.js: file selection, async job polling, result rendering, evidence queries and previews.", "styles.css: responsive workspace, status states, artifact scores and mobile layout.", "Every number is rendered from a backend response, not hardcoded UI state."], 5.35, 2.45, 6.5, 3.2, 16)
code_block(slide, "const job = await fetch('/api/jobs', {\n  method: 'POST',\n  headers: { 'X-Filename': file.name },\n  body: await file.arrayBuffer()\n});", 5.35, 5.55, 6.6, 0.9)

# 5
slide = base("Backend: Python analysis and evidence APIs", "04 / Backend")
text(slide, "backend/server.py", 0.8, 1.8, 3.5, 0.35, 21, GREEN, True, "Aptos Display")
code_block(slide, "POST /api/jobs\nGET  /api/jobs/{jobId}\nPOST /api/ask\nGET  /api/cases/{caseId}/report\nGET  /api/reconstructed/{caseId}/{artifactId}\nGET  /api/viewable/{caseId}/{artifactId}", 0.8, 2.35, 4.0, 2.1)
bullets(slide, ["ThreadingHTTPServer serves the UI and API.", "SQLite stores cases, results and provenance.", "backend/evidence stores input copies by case ID.", "backend/recovered stores raw, reconstructed and normalized outputs."], 5.2, 2.05, 6.6, 2.1, 16)
box(slide, 5.2, 4.65, 6.6, 1.3, MINT, MINT, True)
text(slide, "Safety boundary", 5.5, 4.88, 2, 0.25, 11, GREEN, True)
text(slide, "The source is treated as read-only. CyberTrace analyzes a copy and records SHA-256 provenance before producing output files.", 5.5, 5.2, 5.7, 0.55, 14, NAVY)

# 6
slide = base("What happens after a fragmented file is uploaded", "05 / Runtime flow", True)
flow = [("01", "Acquire", "Read evidence copy + hash"), ("02", "Extract", "Find headers, footers and gaps"), ("03", "Reconstruct", "Carve raw and best-effort output"), ("04", "Validate", "Decode image / parse PDF"), ("05", "Rank", "Integrity + ML recovery score"), ("06", "Explain", "Graph, report and sources")]
for i, (num, title, detail) in enumerate(flow):
    x = 0.75 + (i % 3) * 4.15; y = 1.8 + (i // 3) * 2.25
    box(slide, x, y, 3.45, 1.55, WHITE, WHITE, True)
    text(slide, num, x + 0.2, y + 0.2, 0.45, 0.3, 12, GREEN, True)
    text(slide, title, x + 0.8, y + 0.18, 2.3, 0.3, 18, NAVY, True, "Aptos Display")
    text(slide, detail, x + 0.8, y + 0.65, 2.3, 0.5, 12, MUTED)
text(slide, "The progress bar reflects actual backend stages from the analysis job.", 0.8, 6.75, 7, 0.3, 12, RGBColor(166,199,188))

# 7
slide = base("Fragment intelligence and integrity", "06 / Analysis models")
box(slide, 0.8, 1.8, 5.7, 4.65, WHITE, RGBColor(220,228,224), True)
text(slide, "Fragment relationship graph", 1.1, 2.1, 4.7, 0.35, 20, GREEN, True, "Aptos Display")
text(slide, "Nodes represent recovered artifacts and unresolved clusters. Edges encode adjacency or cluster containment.", 1.1, 2.65, 4.7, 0.7, 15, MUTED)
box(slide, 1.2, 3.65, 1.35, 0.55, MINT, MINT, True); text(slide, "artifact-1", 1.2, 3.65, 1.35, 0.55, 11, GREEN, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
box(slide, 3.75, 3.65, 1.8, 0.55, RGBColor(247,234,214), RGBColor(247,234,214), True); text(slide, "cluster · 32 B", 3.75, 3.65, 1.8, 0.55, 11, ORANGE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
line = slide.shapes.add_connector(1, Inches(2.55), Inches(3.92), Inches(3.75), Inches(3.92)); line.line.color.rgb = GREEN
text(slide, "contains unresolved cluster", 2.5, 4.35, 2.7, 0.25, 10, MUTED, align=PP_ALIGN.CENTER)
box(slide, 6.9, 1.8, 5.7, 4.65, MINT, MINT, True)
text(slide, "Integrity model", 7.2, 2.1, 3, 0.35, 20, GREEN, True, "Aptos Display")
bullets(slide, ["Shannon entropy profile", "Header/footer validation", "Missing-cluster ratio", "Completeness and confidence", "Format-aware image/PDF validation"], 7.2, 2.7, 4.8, 2.4, 16)
text(slide, "ML rank features", 7.2, 5.35, 2, 0.25, 11, GREEN, True)
text(slide, "integrity · confidence · footer · gap coverage · priority · entropy", 7.2, 5.68, 4.7, 0.4, 13, NAVY)

# 8
slide = base("ML ranking and investigator decisions", "07 / Intelligence")
text(slide, "RecoveryRanker", 0.85, 1.85, 4, 0.35, 22, GREEN, True, "Aptos Display")
code_block(slide, "features = {\n  integrity, confidence, footer,\n  gap_free, priority, entropy\n}\nml_score = ranker.predict(features)", 0.85, 2.45, 4.4, 1.65)
bullets(slide, ["Dependency-free logistic regression trained at startup.", "Ranks artifacts for restoration priority.", "Decision labels make the score actionable.", "The model is transparent and replaceable with a larger trained model."], 6.0, 2.05, 5.8, 2.3, 17)
for i, (label, color, score) in enumerate([("RESTORE FIRST", GREEN, "80-100"), ("RESTORE AFTER REVIEW", ORANGE, "55-79"), ("MANUAL INVESTIGATION", RED, "0-54")]):
    y = 4.8 + i * 0.52
    box(slide, 6.0, y, 2.4, 0.35, color, color, True); text(slide, label, 6.0, y, 2.4, 0.35, 9, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, score, 8.65, y, 1.2, 0.35, 13, NAVY, True)

# 9
slide = base("What investigators see", "08 / Product experience", True)
items = [("Evidence queue", "Classified artifacts, integrity, confidence and ML rank"), ("Artifact inspector", "Supported preview, reconstruction sizes and evidence source details"), ("Relationships", "Artifact and unresolved-cluster graph"), ("Ask CyberTrace", "Local evidence retrieval with source artifacts"), ("Reports", "Case ID, input hash, artifacts, graph and recommendations")]
for i, (title, detail) in enumerate(items):
    y = 1.75 + i * 0.95
    box(slide, 1.0, y, 0.5, 0.5, GREEN, GREEN, True); text(slide, str(i+1).zfill(2), 1.0, y, 0.5, 0.5, 11, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, title, 1.8, y + 0.02, 2.6, 0.25, 17, WHITE, True)
    text(slide, detail, 4.35, y + 0.04, 7.5, 0.35, 14, RGBColor(188,211,204))

# 10
slide = base("Storage, provenance and outputs", "09 / Evidence handling")
paths = [("backend/evidence/", "Stored copy of each uploaded input", GREEN), ("backend/recovery_cases.sqlite3", "Case metadata, JSON results, hashes and provenance", NAVY), ("backend/recovered/", "Raw carved, reconstructed and normalized outputs", ORANGE)]
for i, (path, detail, color) in enumerate(paths):
    y = 1.9 + i * 1.35
    box(slide, 0.9, y, 3.6, 0.85, color, color, True)
    text(slide, path, 1.15, y + 0.17, 3.1, 0.25, 16, WHITE, True, "Consolas")
    text(slide, detail, 5.0, y + 0.15, 6.7, 0.4, 17, NAVY)
text(slide, "Input provenance", 0.9, 6.05, 2.2, 0.3, 14, GREEN, True)
text(slide, "SHA-256 + byte count + timestamp + read-only acquisition event", 3.0, 6.05, 7.5, 0.3, 16, MUTED)

# 11
slide = base("Demonstrated results", "10 / Verification")
box(slide, 0.85, 1.85, 3.1, 3.8, NAVY, NAVY, True)
text(slide, "LIVE CHECKS", 1.15, 2.15, 2.2, 0.3, 11, MINT, True)
for i, line in enumerate(["JPEG reconstruction: valid", "PDF parser: 1 page", "ML model: logistic regression", "Graph: nodes + edges", "PNG output: image/png", "Case persistence: SQLite"]):
    text(slide, "✓  " + line, 1.15, 2.75 + i * 0.42, 2.5, 0.25, 12, WHITE)
text(slide, "Evaluation corpus", 4.7, 1.95, 3.5, 0.35, 22, GREEN, True, "Aptos Display")
text(slide, "12 generated files across four formats and three conditions.", 4.7, 2.55, 5.5, 0.35, 16, MUTED)
for i, (label, value) in enumerate([("Formats", "JPEG / PDF / SQLite / ZIP"), ("Conditions", "Intact / fragmented / corrupted"), ("Type classification", "12 / 12 in local check"), ("Known gap", "ZIP and SQLite validators next")]):
    y = 3.25 + i * 0.62
    text(slide, label, 4.7, y, 2.0, 0.25, 12, MUTED, True)
    text(slide, value, 7.0, y, 4.7, 0.3, 14, NAVY)

# 12
slide = base("Next engineering steps", "11 / Roadmap", True)
bullets(slide, ["Filesystem-aware deleted-file recovery: NTFS, FAT32, exFAT and unallocated clusters.", "Streaming disk-image processing for large .img and .dd evidence.", "SQLite integrity checks and ZIP CRC validation.", "Larger labeled corpus with precision, recall and reconstruction metrics.", "Optional embeddings/RAG/LLM layer with mandatory evidence citations.", "Production safeguards: authentication, HTTPS, isolated storage and chain-of-custody controls."], 1.0, 1.8, 11.2, 4.5, 18, WHITE)
box(slide, 1.0, 6.05, 8.9, 0.65, GREEN, GREEN, True)
text(slide, "Current claim: working AI-assisted recovery application with transparent limits.", 1.0, 6.05, 8.9, 0.65, 14, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print(OUT)
